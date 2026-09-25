#!/usr/bin/env python3
"""
Scribe — локальная транскрипция аудио и видео в текст на MLX Whisper (Apple Silicon).

Интерфейс на customtkinter. Возможности:
  • перетаскивание файлов в окно (drag-and-drop) или выбор кнопкой
  • язык: авто / русский / английский
  • модель: Turbo (быстрая, работает локально)
  • формат результата: Текст (.txt), Word (.docx) или Субтитры (.srt)
  • разбивка текста на абзацы; опциональные таймкоды
  • опционально: определение говорящих (диаризация через pyannote)
  • настоящий прогресс-бар и таймер

Результат сохраняется рядом с исходным аудиофайлом.
"""

import os
import io
import re
import json
import time
import queue
import threading
import subprocess
import contextlib
from pathlib import Path
import sys
import platform

# На Apple Silicon MLX работает только в arm64. Если приложение запустили в
# x86_64 (Rosetta или .app с чужой архитектурой) — перезапускаем себя нативно,
# иначе mlx.core не загрузится (ошибка "incompatible architecture").
if (sys.platform == "darwin" and platform.machine() == "x86_64"
        and not os.environ.get("SCRIBE_ARM64")):
    os.environ["SCRIBE_ARM64"] = "1"
    try:
        os.execvp("arch", ["arch", "-arm64", sys.executable] + sys.argv)
    except Exception:
        pass

# В собранном .app (py2app) главный исполняемый файл — сам Scribe, а не python.
# Поэтому процесс распознавания диктофона запускается перезапуском приложения
# с этим флагом. Ловим его до поднятия интерфейса и уходим в воркер.
if len(sys.argv) >= 2 and sys.argv[1] == "--dictaphone-worker":
    import dictaphone_worker
    sys.argv = [sys.argv[0]] + sys.argv[2:]
    dictaphone_worker.main()
    sys.exit(0)

# --- PATH: чтобы ffmpeg/ffprobe нашлись даже при запуске из Finder ---
for _p in ("/opt/homebrew/bin", "/usr/local/bin", os.path.expanduser("~/.local/bin")):
    if os.path.isdir(_p) and _p not in os.environ.get("PATH", "").split(os.pathsep):
        os.environ["PATH"] = _p + os.pathsep + os.environ.get("PATH", "")

try:
    import customtkinter as ctk
    from tkinter import filedialog, messagebox
except Exception:
    import tkinter as tk
    from tkinter import messagebox
    r = tk.Tk(); r.withdraw()
    messagebox.showerror(
        "Нужен модуль customtkinter",
        "Установите:\n\n/Library/Frameworks/Python.framework/Versions/3.14/bin/"
        "python3 -m pip install customtkinter")
    raise SystemExit(1)

HAS_DND = False
try:
    from tkinterdnd2 import TkinterDnD, DND_FILES

    class RootTk(ctk.CTk, TkinterDnD.DnDWrapper):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self.TkdndVersion = TkinterDnD._require(self)

    HAS_DND = True
except Exception:
    RootTk = ctk.CTk


# ---------------------------------------------------------------------------
# Turbo — модель по умолчанию (быстрая). Large v3 доступна в Настройках как
# опция «на всякий», но в разы медленнее при том же результате.
MODEL_REPO = "mlx-community/whisper-large-v3-turbo"
MODELS = {
    "Turbo — быстро (рекомендуется)": "mlx-community/whisper-large-v3-turbo",
    "Large v3 — точнее, но в разы медленнее": "mlx-community/whisper-large-v3-mlx",
}
# Язык — сегменты; «Несколько» включает блочный мультиязычный движок.
LANG_SEG = ["Авто", "Русский", "English", "Несколько"]
LANG_CODE = {"Русский": "ru", "English": "en"}   # «Авто»/«Несколько» → None
FMT_SEG = {"Текст": "txt", "Word": "docx", "Субтитры": "srt"}
NSPK = {"Определить само": 0, "2": 2, "3": 3, "4": 4, "5": 5, "6": 6}


def _model_title(repo):
    for k, v in MODELS.items():
        if v == repo:
            return k
    return list(MODELS)[0]


def _models_base(cfg):
    return cfg.get("models_dir") or os.environ.get("HF_HOME") \
        or os.path.expanduser("~/.cache/huggingface")


def _models_dir_path(cfg):
    return os.path.join(_models_base(cfg), "hub")


def _models_dir_display(cfg):
    return _models_dir_path(cfg).replace(os.path.expanduser("~"), "~")


def _apply_models_dir(d):
    if d:
        os.environ["HF_HOME"] = d

AUDIO_EXTS = {
    # аудио
    ".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg", ".oga", ".opus",
    ".aiff", ".aif", ".wma", ".m4b", ".amr", ".mka", ".ape", ".wv", ".caf", ".ac3",
    # видео
    ".mp4", ".mov", ".m4v", ".mkv", ".webm", ".wmv", ".avi", ".flv",
    ".3gp", ".3g2", ".ts", ".mts", ".m2ts", ".mpg", ".mpeg", ".asf", ".vob", ".ogv",
}
AUDIO_TYPES = [("Аудио и видео", " ".join("*" + e for e in sorted(AUDIO_EXTS))),
               ("Все файлы", "*.*")]

# Список поддерживаемых форматов — показываем в «Настройках».
FORMATS_AUDIO = "MP3, WAV, M4A, AAC, FLAC, OGG, OPUS, AIFF, WMA, M4B, AMR"
FORMATS_VIDEO = "MP4, MOV, MKV, WEBM, WMV, AVI, FLV, 3GP, TS, MPG"

APP_VERSION = "2.0"   # внутренняя версия, в интерфейсе не показывается
CONFIG_PATH = Path(os.path.expanduser("~/.audio2text.json"))
# Пробуем модели по порядку: сначала новая community-1 (её тянут свежие версии
# pyannote и она точнее), затем классическая 3.1 как запасная.
DIAR_MODELS = [
    "pyannote/speaker-diarization-community-1",
    "pyannote/speaker-diarization-3.1",
]
# Что проверять на доступ: community-1 самодостаточна; 3.1 требует ещё segmentation-3.0.
DIAR_CHECK = [
    "pyannote/speaker-diarization-community-1",
    "pyannote/speaker-diarization-3.1",
    "pyannote/segmentation-3.0",
]


def load_config() -> dict:
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(cfg: dict) -> None:
    try:
        CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


# ---------------------------------------------------------------------------
def parse_ts(text: str) -> float:
    try:
        parts = [float(p) for p in text.split(":")]
    except ValueError:
        return 0.0
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    return parts[0]


def ffprobe_duration(path) -> float:
    try:
        out = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
            capture_output=True, text=True)
        return float(out.stdout.strip())
    except Exception:
        return 0.0


def _clock(seconds: float, srt: bool = False) -> str:
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    if srt:
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def group_paragraphs(segments, gap=1.2, max_chars=380):
    paras, buf, start, prev_end = [], [], None, None
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        if start is None:
            start = seg["start"]
        big_pause = prev_end is not None and (seg["start"] - prev_end) > gap
        cur = " ".join(buf)
        long_enough = len(cur) > max_chars and cur.rstrip().endswith((".", "!", "?", "…"))
        if buf and (big_pause or long_enough):
            paras.append((start, " ".join(buf).strip()))
            buf, start = [], seg["start"]
        buf.append(text)
        prev_end = seg["end"]
    if buf:
        paras.append((start, " ".join(buf).strip()))
    return paras


def group_by_speaker(segments, gap=1.2, max_chars=380):
    """Блоки (start, speaker, text): новый блок при смене говорящего,
    большой паузе или когда абзац стал длинным."""
    blocks, buf, start, cur, prev_end = [], [], None, None, None
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        spk = seg.get("speaker", "Спикер ?")
        if cur is None:
            cur, start = spk, seg["start"]
        speaker_changed = spk != cur
        big_pause = prev_end is not None and (seg["start"] - prev_end) > gap
        cur_text = " ".join(buf)
        long_enough = len(cur_text) > max_chars and cur_text.rstrip().endswith((".", "!", "?", "…"))
        if buf and (speaker_changed or big_pause or long_enough):
            blocks.append((start, cur, " ".join(buf).strip()))
            buf, start, cur = [], seg["start"], spk
        buf.append(text)
        prev_end = seg["end"]
    if buf:
        blocks.append((start, cur, " ".join(buf).strip()))
    return blocks


def write_txt(segments, path, with_ts, diarized):
    lines = []
    if diarized:
        prev_spk = None
        for start, spk, text in group_by_speaker(segments):
            ts = f"[{_clock(start)}] " if with_ts else ""
            if spk != prev_spk:
                lines.append(f"{ts}{spk}: {text}")
            else:
                lines.append(f"{ts}{text}")
            lines.append("")
            prev_spk = spk
    else:
        for start, text in group_paragraphs(segments):
            p = f"[{_clock(start)}] " if with_ts else ""
            lines.append(f"{p}{text}"); lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def write_srt(segments, path, diarized):
    blocks = []
    for i, seg in enumerate(segments, 1):
        text = seg["text"].strip()
        if diarized and seg.get("speaker"):
            text = f"{seg['speaker']}: {text}"
        blocks.append(str(i))
        blocks.append(f"{_clock(seg['start'], True)} --> {_clock(seg['end'], True)}")
        blocks.append(text); blocks.append("")
    path.write_text("\n".join(blocks), encoding="utf-8")


def write_docx(segments, path, title, with_ts, diarized):
    from docx import Document
    from docx.shared import RGBColor
    doc = Document()
    doc.add_heading(title, level=1)
    if diarized:
        prev_spk = None
        for start, spk, text in group_by_speaker(segments):
            p = doc.add_paragraph()
            if with_ts:
                ts = p.add_run(f"[{_clock(start)}] "); ts.italic = True
                ts.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
            if spk != prev_spk:
                r = p.add_run(f"{spk}: "); r.bold = True
                r.font.color.rgb = RGBColor(0x2f, 0x6f, 0xed)
            p.add_run(text)
            prev_spk = spk
    else:
        for start, text in group_paragraphs(segments):
            p = doc.add_paragraph()
            if with_ts:
                t = p.add_run(f"[{_clock(start)}] "); t.italic = True
                t.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
            p.add_run(text)
    doc.save(str(path))


# ---------------------------------------------------------------------------
def detect_silences(path, noise="-30dB", min_dur=0.5):
    """Находит паузы через ffmpeg silencedetect -> [(start, end), ...]."""
    out = subprocess.run(
        ["ffmpeg", "-i", str(path), "-af",
         f"silencedetect=noise={noise}:d={min_dur}", "-f", "null", "-"],
        capture_output=True, text=True)
    txt = out.stderr
    starts = [float(x) for x in re.findall(r"silence_start:\s*([\d.]+)", txt)]
    ends = [float(x) for x in re.findall(r"silence_end:\s*([\d.]+)", txt)]
    return list(zip(starts, ends))


def build_chunks(path, target=45.0, min_gap=6.0):
    """Режет запись на куски ~target секунд, разрезая только по паузам."""
    dur = ffprobe_duration(path)
    if dur <= 0:
        return [(0.0, 0.0)]
    mids = [(s + e) / 2 for s, e in detect_silences(path)]
    cuts = [0.0]
    tt = target
    while tt < dur:
        near = [m for m in mids if m > cuts[-1] + min_gap]
        if near:
            best = min(near, key=lambda m: abs(m - tt))
            if best < dur - 2:
                cuts.append(best)
        tt += target
    cuts.append(dur)
    chunks = [(cuts[i], cuts[i + 1]) for i in range(len(cuts) - 1)
              if cuts[i + 1] - cuts[i] > 0.3]
    return chunks or [(0.0, dur)]


def transcribe_multilang(src, model, mlx_whisper, prog, log):
    """Режет запись на куски по паузам и распознаёт каждый со своим языком (авто)."""
    import tempfile
    chunks = build_chunks(src)
    total = len(chunks)
    all_seg = []
    for i, (a, b) in enumerate(chunks):
        log(f"Кусок {i + 1} из {total} (язык — авто)…")
        clip = Path(tempfile.gettempdir()) / f"_ml_{i}.wav"
        subprocess.run(
            ["ffmpeg", "-ss", str(a), "-i", str(src), "-t", str(b - a),
             "-ar", "16000", "-ac", "1", "-c:a", "pcm_s16le", str(clip),
             "-y", "-hide_banner", "-loglevel", "error"], check=True)
        try:
            res = mlx_whisper.transcribe(
                str(clip), path_or_hf_repo=model, language=None,
                temperature=0.0, condition_on_previous_text=False, verbose=False)
            for seg in res.get("segments", []):
                seg["start"] += a
                seg["end"] += a
                all_seg.append(seg)
        finally:
            try: clip.unlink()
            except Exception: pass
        prog((i + 1) / total)
    return all_seg


# ---------------------------------------------------------------------------
def normalize_audio(path, log):
    """Выравнивает громкость через ffmpeg loudnorm во временный wav 16 кГц."""
    import tempfile
    tmp = Path(tempfile.gettempdir()) / (path.stem + "_norm.wav")
    log("Нормализую громкость…")
    subprocess.run(
        ["ffmpeg", "-i", str(path), "-af", "loudnorm", "-ar", "16000", "-ac", "1",
         "-c:a", "pcm_s16le", str(tmp), "-y", "-hide_banner", "-loglevel", "error"],
        check=True)
    return tmp


# ---------------------------------------------------------------------------
# Быстрый мультиязычный движок (Scribe v2): один декод аудио в память, дешёвое
# разреженное определение языка по записи, склейка соседних кусков одного языка
# в крупные блоки, транскрипция каждого блока ОДНИМ вызовом с явным языком.
# Даёт и скорость (меньше проходов), и чистые границы языков (нет дрейфа).
SR_FAST = 16000


def _decode_to_array(path, normalize=True):
    """ОДИН проход ffmpeg -> float32 моно 16 кГц в память. loudnorm вшит сюда же,
    так что отдельного прохода нормализации не нужно."""
    import numpy as np
    cmd = ["ffmpeg", "-nostdin", "-i", str(path)]
    if normalize:
        cmd += ["-af", "loudnorm"]
    cmd += ["-ar", str(SR_FAST), "-ac", "1", "-f", "f32le", "-acodec", "pcm_f32le",
            "-hide_banner", "-loglevel", "error", "-"]
    proc = subprocess.run(cmd, capture_output=True)
    if proc.returncode != 0:
        raise RuntimeError("ffmpeg: " + proc.stderr.decode("utf-8", "ignore")[-300:])
    return np.frombuffer(proc.stdout, dtype=np.float32).copy()


def _find_silences_arr(audio, sr=SR_FAST, noise_db=-30.0, min_dur=0.5, frame=0.03, hop=0.01):
    """Паузы [(start, end), ...] по массиву в памяти (без второго прохода ffmpeg)."""
    import numpy as np
    fl = max(1, int(frame * sr)); hl = max(1, int(hop * sr))
    if len(audio) < fl:
        return []
    n = 1 + (len(audio) - fl) // hl
    idx = np.arange(fl)[None, :] + hl * np.arange(n)[:, None]
    rms = np.sqrt(np.mean(audio[idx].astype(np.float64) ** 2, axis=1) + 1e-12)
    db = 20.0 * np.log10(rms / max(rms.max(), 1e-9))
    quiet = db < noise_db
    sil, i = [], 0
    while i < n:
        if quiet[i]:
            j = i
            while j < n and quiet[j]:
                j += 1
            s, e = i * hop, (j - 1) * hop + frame
            if e - s >= min_dur:
                sil.append((s, e))
            i = j
        else:
            i += 1
    return sil


class LangDetector:
    """Грузит модель один раз; дёшево определяет язык 30-секундного окна. Сам
    подбирает рабочую форму мел-спектра под конкретную версию mlx_whisper и
    запоминает её (у разных версий API немного отличается)."""

    def __init__(self, model_repo):
        import mlx.core as mx
        import mlx_whisper
        from mlx_whisper import audio as mlx_audio
        self.mx = mx
        self.mlx_whisper = mlx_whisper
        self.mlx_audio = mlx_audio
        self.model_repo = model_repo
        self.calls = 0
        self._calib = None
        self._detect_fn = None
        self.method = None
        self._load_model()
        self._pick_method()

    def _load_model(self):
        mx = self.mx
        try:
            from mlx_whisper.transcribe import ModelHolder
            self.model = ModelHolder.get_model(self.model_repo, mx.float16); return
        except Exception:
            pass
        from mlx_whisper.load_models import load_model
        self.model = load_model(self.model_repo)

    def _pick_method(self):
        try:
            from mlx_whisper.decoding import detect_language as fn
            self._detect_fn = fn; self.method = "decoding"; return
        except Exception:
            pass
        self.method = "fallback"

    @property
    def n_mels(self):
        try:
            return int(self.model.dims.n_mels)
        except Exception:
            return 128

    def _mel(self, chunk):
        from mlx_whisper.audio import pad_or_trim, N_FRAMES
        mel = self.mlx_audio.log_mel_spectrogram(chunk, self.n_mels)
        if mel.shape[-1] == self.n_mels:
            mel = pad_or_trim(mel, N_FRAMES, axis=0)
        elif mel.shape[0] == self.n_mels:
            mel = pad_or_trim(mel, N_FRAMES, axis=-1)
        return mel.astype(self.mx.float16)

    def _variants(self, mel):
        out = [("2d", mel), ("3d", mel[None])]
        try:
            out += [("t2d", mel.T), ("t3d", mel.T[None])]
        except Exception:
            pass
        return out

    def _apis(self):
        a = []
        if self._detect_fn is not None:
            a.append(lambda m: self._detect_fn(self.model, m))
        if hasattr(self.model, "detect_language"):
            a.append(lambda m: self.model.detect_language(m))
        return a

    @staticmethod
    def _parse(out):
        probs = out[1] if isinstance(out, (tuple, list)) and len(out) >= 2 else out
        if isinstance(probs, list):
            probs = probs[0]
        if isinstance(probs, dict):
            return max(probs, key=probs.get)
        return str(probs)

    def detect(self, chunk):
        self.calls += 1
        if self.method == "fallback":
            res = self.mlx_whisper.transcribe(
                chunk, path_or_hf_repo=self.model_repo, language=None,
                temperature=0.0, condition_on_previous_text=False, verbose=False)
            return res.get("language", "??")
        mel = self._mel(chunk)
        variants = dict(self._variants(mel))
        if self._calib is not None:
            afn, vname = self._calib
            return self._parse(afn(variants[vname]))
        for afn in self._apis():
            for vname, m in self._variants(mel):
                try:
                    lang = self._parse(afn(m))
                    self._calib = (afn, vname)
                    return lang
                except Exception:
                    continue
        raise RuntimeError("Не удалось определить язык (detect_language).")


def _detect_window(det, audio, center_s, win=30.0):
    b = int(min(len(audio), (center_s + win / 2) * SR_FAST))
    a = int(max(0, (center_s - win / 2) * SR_FAST))
    chunk = audio[a:b]
    if len(chunk) < int(1.0 * SR_FAST):
        chunk = audio[max(0, b - int(win * SR_FAST)):b]
    return det.detect(chunk)


def _smart_scan(det, audio, silences, coarse=150.0, precision=20.0,
                dense_below=360.0, prog=None):
    """Разреженные якоря + бисекция границ языков. Короткие файлы (< dense_below)
    сканируются плотно — там это дёшево. -> [(start, end, lang), ...]."""
    import numpy as np
    total = len(audio) / SR_FAST
    sil_mids = [(s + e) / 2 for s, e in silences]
    step = 30.0 if total <= dense_below else coarse
    grid = [i * step + step / 2 for i in range(max(1, int(np.ceil(total / step))))]
    edge = min(15.0, total / 2.0)   # обязательные якоря у краёв: не терять интро/концовку
    anchors = sorted({min(total - 0.5, max(0.0, a)) for a in ([edge] + grid + [total - edge])})
    labels = []
    for c in anchors:
        labels.append((c, _detect_window(det, audio, c)))
        if prog:
            prog(len(labels) / max(1, len(anchors)))
    blocks, start, cur = [], 0.0, labels[0][1]
    for i in range(len(labels) - 1):
        (ca, la), (cb, lb) = labels[i], labels[i + 1]
        if la == lb:
            continue
        lo, hi, llo = ca, cb, la
        while hi - lo > precision:
            mid = (lo + hi) / 2
            if _detect_window(det, audio, mid) == llo:
                lo = mid
            else:
                hi = mid
        bt = (lo + hi) / 2
        near = [m for m in sil_mids if abs(m - bt) <= 8.0]
        if near:
            bt = min(near, key=lambda m: abs(m - bt))
        blocks.append((start, bt, cur)); start, cur = bt, lb
    blocks.append((start, total, cur))
    return [(a, b, l) for a, b, l in blocks if b - a > 0.2]


def transcribe_multilang_fast(path, model_repo, mlx_whisper, normalize, prog, stat):
    """Новый мультиязычный движок. Возвращает segments (start/end/text)."""
    import numpy as np
    stat("Готовлю аудио…")
    audio = _decode_to_array(path, normalize)
    total_dur = len(audio) / SR_FAST if len(audio) else 0.0
    stat("Ищу паузы…")
    sil = _find_silences_arr(audio)
    stat("Определяю языки по записи…")
    det = LangDetector(model_repo)
    blocks = _smart_scan(det, audio, sil, prog=lambda fr: prog(0.15 * fr))
    segments, done = [], 0.0
    for k, (a, b, lang) in enumerate(blocks, 1):
        stat(f"Распознаю блок {k} из {len(blocks)} · {lang}…")
        clip = np.ascontiguousarray(audio[int(a * SR_FAST):int(b * SR_FAST)])
        res = mlx_whisper.transcribe(
            clip, path_or_hf_repo=model_repo, language=lang,
            temperature=0.0, condition_on_previous_text=False, verbose=False)
        for seg in res.get("segments", []):
            seg = dict(seg); seg["start"] += a; seg["end"] += a
            segments.append(seg)
        done += (b - a)
        prog(0.15 + 0.85 * (done / total_dur if total_dur else 1.0))
    return segments


# ---------------------------------------------------------------------------
def _load_pipeline(token):
    """Загружает pyannote pipeline независимо от версии: логинимся заранее,
    токен подхватывается из кэша/окружения, а не через параметр (у разных
    версий pyannote параметр называется по-разному: token / use_auth_token)."""
    os.environ["HF_TOKEN"] = token
    os.environ["HUGGING_FACE_HUB_TOKEN"] = token
    try:
        from huggingface_hub import login
        login(token=token, add_to_git_credential=False)
    except Exception:
        pass
    from pyannote.audio import Pipeline
    errors = []
    for model in DIAR_MODELS:
        for kw in ({}, {"token": token}):
            try:
                pl = Pipeline.from_pretrained(model, **kw)
                if pl is not None:
                    return pl
            except Exception as e:
                errors.append(f"{model}: {e}")
    raise RuntimeError(
        "Ни одна модель диаризации не загрузилась. Обычно причина — не приняты "
        "условия модели на Hugging Face. Откройте и нажмите Agree / Accept на:\n"
        "huggingface.co/pyannote/speaker-diarization-community-1\n\n"
        + "\n".join(errors[:4]))


def _mk_hook(prog, stat):
    """Колбэк прогресса pyannote: превращает шаги диаризации в прогресс-бар."""
    def hook(step_name, step_artifact=None, file=None, total=None, completed=None):
        try:
            if total and completed is not None and float(total) > 0:
                prog(min(1.0, float(completed) / float(total)))
            if step_name:
                stat(f"Говорящие: {step_name}…")
        except Exception:
            pass
    return hook


def _extract_turns(diar):
    """Достаёт (start, end, speaker) из результата любой версии pyannote."""
    ann = getattr(diar, "speaker_diarization", diar)
    turns = []
    if hasattr(ann, "itertracks"):
        for item in ann.itertracks(yield_label=True):
            seg, spk = item[0], item[-1]
            turns.append((seg.start, seg.end, spk))
    else:
        for turn, spk in ann:
            turns.append((turn.start, turn.end, spk))
    return turns


def diarize_and_label(audio_path, segments, token, num_speakers, log, prog):
    import torch
    log("Загружаю модель диаризации…")
    pipeline = _load_pipeline(token)
    device = "cpu"
    try:
        if torch.backends.mps.is_available():
            pipeline.to(torch.device("mps")); device = "mps"
    except Exception:
        device = "cpu"
    # Готовим аудио сами через ffmpeg и подаём waveform напрямую — минуя torchcodec,
    # который ломается на свежих связках PyTorch/ffmpeg.
    import wave, tempfile
    import numpy as np
    log("Готовлю аудио для диаризации…")
    tmp = Path(tempfile.gettempdir()) / (audio_path.stem + "_diar.wav")
    subprocess.run(
        ["ffmpeg", "-i", str(audio_path), "-ar", "16000", "-ac", "1",
         "-c:a", "pcm_s16le", str(tmp), "-y", "-hide_banner", "-loglevel", "error"],
        check=True)
    hook = _mk_hook(prog, log)

    def _run(pl):
        try:
            return pl(payload, hook=hook, **kwargs)
        except TypeError:
            return pl(payload, **kwargs)  # старые версии без hook

    try:
        with wave.open(str(tmp), "rb") as w:
            sr = w.getframerate()
            frames = w.readframes(w.getnframes())
        data = np.frombuffer(frames, dtype=np.int16).astype(np.float32) / 32768.0
        waveform = torch.from_numpy(data).unsqueeze(0)  # (1, N)
        payload = {"waveform": waveform, "sample_rate": sr}
        log(f"Определяю говорящих (устройство: {device})…")
        kwargs = {"num_speakers": num_speakers} if num_speakers and num_speakers > 0 else {}
        try:
            diar = _run(pipeline)
        except Exception as e:
            if device == "mps":
                log("MPS не справился, повтор на CPU…")
                pipeline.to(torch.device("cpu"))
                diar = _run(pipeline)
            else:
                raise e
        turns = _extract_turns(diar)
    finally:
        try: tmp.unlink()
        except Exception: pass
    mapping, n = {}, 0
    for _, _, spk in turns:
        if spk not in mapping:
            n += 1; mapping[spk] = f"Спикер {n}"

    def speaker_at(a, b):
        best, best_ov = None, 0.0
        for s, e, spk in turns:
            ov = max(0.0, min(b, e) - max(a, s))
            if ov > best_ov:
                best_ov, best = ov, spk
        return mapping.get(best, "Спикер ?")

    for seg in segments:
        seg["speaker"] = speaker_at(seg["start"], seg["end"])
    return segments


# ---------------------------------------------------------------------------
class SettingsWindow(ctk.CTkToplevel):
    def __init__(self, master, app):
        super().__init__(master)
        self.app = app
        self.title("Настройки")
        self.geometry("560x620")
        self.transient(master)
        P = {"padx": 20}

        # --- Модель ---
        ctk.CTkLabel(self, text="Модель распознавания",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", pady=(18, 4), **P)
        self.model_var = ctk.StringVar(value=_model_title(app.cfg.get("model_repo", MODEL_REPO)))
        ctk.CTkOptionMenu(self, variable=self.model_var, values=list(MODELS),
                          width=440, command=self._save_model).pack(anchor="w", **P)
        ctk.CTkLabel(self, text="Turbo — быстрая и рекомендуемая. Large v3 точнее в теории, "
                                "но в разы медленнее при том же результате.",
                     text_color="gray", justify="left", wraplength=500).pack(anchor="w", pady=(3, 0), **P)

        # --- Папка моделей ---
        ctk.CTkLabel(self, text="Папка моделей",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", pady=(16, 4), **P)
        mrow = ctk.CTkFrame(self, fg_color="transparent"); mrow.pack(fill="x", **P)
        self.models_dir_lbl = ctk.CTkLabel(mrow, text=_models_dir_display(app.cfg),
                                           text_color="gray", anchor="w")
        self.models_dir_lbl.pack(side="left", fill="x", expand=True)
        ctk.CTkButton(mrow, text="Сменить…", width=90, command=self._change_models).pack(side="right")
        ctk.CTkButton(mrow, text="Открыть", width=90, command=self._open_models).pack(side="right", padx=6)
        ctk.CTkLabel(self, text="Сюда скачиваются модели (~1.5 ГБ). Можно открыть папку и удалить ненужное.",
                     text_color="gray", justify="left", wraplength=500).pack(anchor="w", pady=(3, 0), **P)

        # --- Токен Hugging Face ---
        ctk.CTkLabel(self, text="Токен Hugging Face",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", pady=(16, 2), **P)
        ctk.CTkLabel(self, text="Нужен только для «Определять говорящих». Разовая настройка.",
                     text_color="gray").pack(anchor="w", **P)
        self._saved_token = app.cfg.get("hf_token", "")
        self.entry = ctk.CTkEntry(self, width=500, show="•", placeholder_text="hf_...")
        self.entry.pack(pady=(8, 6), **P); self.entry.insert(0, self._saved_token)
        self.entry.bind("<KeyRelease>", lambda e: self._sync_btn())
        row = ctk.CTkFrame(self, fg_color="transparent"); row.pack(fill="x", **P)
        self.check_btn = ctk.CTkButton(row, text="Проверить и сохранить", width=210, command=self._check)
        self.check_btn.pack(side="left")
        self.status = ctk.CTkLabel(self, text="", justify="left", wraplength=500)
        self.status.pack(anchor="w", pady=(8, 0), **P)

        # --- Поддерживаемые форматы ---
        ctk.CTkLabel(self, text="Поддерживаемые форматы",
                     font=ctk.CTkFont(size=15, weight="bold")).pack(anchor="w", pady=(16, 2), **P)
        ctk.CTkLabel(self, text=f"Аудио: {FORMATS_AUDIO}\nВидео: {FORMATS_VIDEO}\n"
                                "и вообще всё, что умеет открывать ffmpeg.",
                     text_color="gray", justify="left", wraplength=500).pack(anchor="w", pady=(2, 0), **P)

        ctk.CTkLabel(self, text=f"Scribe {APP_VERSION}  ·  локально на Apple MLX",
                     text_color="gray").pack(side="bottom", pady=(0, 12))
        self._sync_btn()

    def _sync_btn(self):
        changed = self.entry.get().strip() != self._saved_token
        if self._saved_token and not changed:
            self.status.configure(text="✓ токен сохранён и проверен", text_color="green")
            self.check_btn.configure(state="disabled")
        else:
            self.check_btn.configure(state="normal")
            if not self._saved_token and not self.entry.get().strip():
                self.status.configure(text="", text_color="gray")

    def _save_model(self, _v=None):
        self.app.cfg["model_repo"] = MODELS[self.model_var.get()]
        save_config(self.app.cfg)

    def _open_models(self):
        p = _models_dir_path(self.app.cfg)
        try: os.makedirs(p, exist_ok=True)
        except Exception: pass
        subprocess.run(["open", p], check=False)

    def _change_models(self):
        d = filedialog.askdirectory(title="Папка для моделей")
        if not d:
            return
        self.app.cfg["models_dir"] = d; save_config(self.app.cfg)
        _apply_models_dir(d)
        self.models_dir_lbl.configure(text=_models_dir_display(self.app.cfg))
        self.status.configure(text="Папка моделей изменена. Модель докачается при первом запуске.",
                              text_color="gray")

    def _check(self):
        token = self.entry.get().strip()
        if not token:
            self.status.configure(text="Введите токен.", text_color="orange"); return
        self.status.configure(text="Проверяю…", text_color="gray"); self.update()
        try:
            from huggingface_hub import HfApi
            api = HfApi(); who = api.whoami(token=token); name = who.get("name", "пользователь")
        except Exception as e:
            self.status.configure(text=f"✗ Токен не принят: {e}", text_color="red"); return
        access = {}
        for m in DIAR_CHECK:
            try:
                api.model_info(m, token=token); access[m] = True
            except Exception:
                access[m] = False
        ok = access.get("pyannote/speaker-diarization-community-1") or (
            access.get("pyannote/speaker-diarization-3.1") and access.get("pyannote/segmentation-3.0"))
        if ok:
            self.app.cfg["hf_token"] = token; save_config(self.app.cfg)
            self._saved_token = token; self.app.refresh_token_badge()
            self.status.configure(text=f"✓ Токен работает ({name}). Сохранён.", text_color="green")
            self._sync_btn()
        else:
            self.status.configure(
                text=f"Токен верный ({name}), но нет доступа к модели диаризации.\n"
                     "Примите условия: huggingface.co/pyannote/speaker-diarization-community-1",
                text_color="orange")


# ---------------------------------------------------------------------------
class App:
    def __init__(self, root):
        self.root = root
        self.cfg = load_config()
        _apply_models_dir(self.cfg.get("models_dir"))
        self.files = []
        self.q = queue.Queue()
        self.running = False
        self.start_time = None
        self.settings_win = None
        self.dict_win = None
        self.last_dir = None

        ctk.set_appearance_mode("system")
        ctk.set_default_color_theme("blue")
        root.title("Scribe")
        root.geometry("660x744")
        root.minsize(620, 700)
        try:
            _icon = Path(__file__).with_name("scribe.png")
            if _icon.exists():
                import tkinter as _tk
                root.iconphoto(True, _tk.PhotoImage(file=str(_icon)))
        except Exception:
            pass

        self._build_ui()
        if HAS_DND:
            root.drop_target_register(DND_FILES)
            root.dnd_bind("<<Drop>>", self._on_drop)
        self.root.after(120, self._drain)
        self.root.after(500, self._tick)

    def _build_ui(self):
        head = ctk.CTkFrame(self.root, fg_color="transparent")
        head.pack(fill="x", padx=20, pady=(14, 4))
        ctk.CTkLabel(head, text="🎙  Scribe",
                     font=ctk.CTkFont(size=24, weight="bold")).pack(side="left")
        ctk.CTkButton(head, text="⚙︎ Настройки", width=110,
                      command=self._open_settings).pack(side="right")
        ctk.CTkButton(head, text="🎤 Диктофон", width=120,
                      command=self._open_dictaphone).pack(side="right", padx=(0, 8))

        body = ctk.CTkFrame(self.root, fg_color="transparent")
        body.pack(fill="x", padx=18, pady=(2, 0))

        self.drop = ctk.CTkFrame(body, height=72, corner_radius=12,
                                 border_width=2, border_color="#8aa0c8")
        self.drop.pack(fill="x", pady=(2, 6)); self.drop.pack_propagate(False)
        hint = "Перетащите сюда аудио или видео" if HAS_DND else "Нажмите, чтобы выбрать файл"
        lbl = ctk.CTkLabel(self.drop, text=f"⬇  {hint}  ·  или нажмите для выбора",
                           font=ctk.CTkFont(size=14), text_color="gray")
        lbl.pack(expand=True)
        for w in (self.drop, lbl):
            w.bind("<Button-1>", lambda e: self.add_files())

        self.listbox = ctk.CTkTextbox(body, height=52, corner_radius=8, border_width=0,
                                      fg_color=("gray92", "gray20"),
                                      text_color=("gray35", "gray72"), activate_scrollbars=False)
        self.listbox.pack(fill="x", pady=(0, 4))
        self._refresh_list()
        btns = ctk.CTkFrame(body, fg_color="transparent"); btns.pack(fill="x", pady=(0, 2))
        ctk.CTkButton(btns, text="Добавить…", width=110, command=self.add_files).pack(side="left")
        ctk.CTkButton(btns, text="Очистить всё", width=110, fg_color="gray",
                      command=self.clear_files).pack(side="left", padx=8)

        # --- Язык (сегменты + подсказка) ---
        lang_card = ctk.CTkFrame(body); lang_card.pack(fill="x", pady=(8, 4))
        lc = ctk.CTkFrame(lang_card, fg_color="transparent"); lc.pack(fill="x", padx=12, pady=(8, 9))
        ctk.CTkLabel(lc, text="ЯЗЫК", font=ctk.CTkFont(size=11, weight="bold"),
                     text_color="gray").pack(anchor="w")
        self.lang_seg = ctk.CTkSegmentedButton(lc, values=LANG_SEG,
                                               command=lambda _v: self._lang_hint())
        self.lang_seg.set(LANG_SEG[0]); self.lang_seg.pack(fill="x", pady=(6, 4))
        self.lang_note = ctk.CTkLabel(lc, text="", font=ctk.CTkFont(size=11),
                                      text_color="gray", anchor="w", justify="left")
        self.lang_note.pack(anchor="w")
        self._lang_hint()

        # --- Формат (сегменты + расшифровка) ---
        fmt_card = ctk.CTkFrame(body); fmt_card.pack(fill="x", pady=4)
        fc = ctk.CTkFrame(fmt_card, fg_color="transparent"); fc.pack(fill="x", padx=12, pady=(8, 9))
        ctk.CTkLabel(fc, text="ФОРМАТ", font=ctk.CTkFont(size=11, weight="bold"),
                     text_color="gray").pack(anchor="w")
        self.fmt_seg = ctk.CTkSegmentedButton(fc, values=list(FMT_SEG))
        self.fmt_seg.set(list(FMT_SEG)[0]); self.fmt_seg.pack(fill="x", pady=(6, 4))
        ctk.CTkLabel(fc, text="Текст = .txt   ·   Word = .docx   ·   Субтитры = .srt",
                     font=ctk.CTkFont(size=11), text_color="gray").pack(anchor="w")

        # --- Опции (тумблеры) ---
        opt = ctk.CTkFrame(body); opt.pack(fill="x", pady=4)
        o = ctk.CTkFrame(opt, fg_color="transparent"); o.pack(fill="x", padx=12, pady=(6, 8))
        r1 = ctk.CTkFrame(o, fg_color="transparent"); r1.pack(fill="x", pady=(2, 0))
        self.diar_var = ctk.BooleanVar(value=True)
        ctk.CTkSwitch(r1, text="Определять говорящих (Спикер 1, 2 … — дольше)",
                      variable=self.diar_var, command=self._toggle_diar).pack(side="left")
        self.token_badge = ctk.CTkLabel(r1, text="", anchor="e"); self.token_badge.pack(side="right")
        r1b = ctk.CTkFrame(o, fg_color="transparent"); r1b.pack(fill="x", pady=(2, 2))
        ctk.CTkLabel(r1b, text="Сколько человек:", text_color="gray").pack(side="left", padx=(46, 8))
        self.nspk_var = ctk.StringVar(value=list(NSPK)[0])
        self.nspk_menu = ctk.CTkOptionMenu(r1b, variable=self.nspk_var, values=list(NSPK), width=160)
        self.nspk_menu.pack(side="left")
        self.norm_var = ctk.BooleanVar(value=True)
        ctk.CTkSwitch(o, text="Нормализовать громкость (меньше ошибок, чуть дольше)",
                      variable=self.norm_var).pack(anchor="w", pady=(12, 8))
        self.ts_var = ctk.BooleanVar(value=False)
        ctk.CTkSwitch(o, text="Добавлять таймкоды (для .txt и Word)",
                      variable=self.ts_var).pack(anchor="w", pady=(0, 4))
        self._toggle_diar(); self.refresh_token_badge()

        # --- Запуск / прогресс / низ ---
        self.run_btn = ctk.CTkButton(self.root, text="Транскрибировать", height=46,
                                     font=ctk.CTkFont(size=16, weight="bold"), command=self.start)
        self.run_btn.pack(fill="x", padx=20, pady=(8, 4))
        bar = ctk.CTkFrame(self.root, fg_color="transparent"); bar.pack(fill="x", padx=20)
        self.progress = ctk.CTkProgressBar(bar); self.progress.set(0)
        self.progress.pack(side="left", fill="x", expand=True)
        self.timer = ctk.CTkLabel(bar, text="⏱ 00:00", width=76); self.timer.pack(side="right", padx=(10, 0))
        foot = ctk.CTkFrame(self.root, fg_color="transparent")
        foot.pack(fill="x", padx=20, pady=(3, 8))
        self.status = ctk.CTkLabel(foot, text="Готово к работе.", anchor="w", text_color="gray")
        self.status.pack(side="left")
        ctk.CTkButton(foot, text="📂 Папка", width=92, fg_color="gray",
                      command=self._open_folder).pack(side="right")

    # -- настройки --
    def _open_settings(self):
        if self.settings_win is None or not self.settings_win.winfo_exists():
            self.settings_win = SettingsWindow(self.root, self)
        self.settings_win.focus()

    def _open_dictaphone(self):
        if self.dict_win is not None and self.dict_win.winfo_exists():
            self.dict_win.focus(); return
        try:
            from dictaphone import DictaphoneWindow
        except Exception as e:
            messagebox.showerror("Диктофон", f"Не удалось загрузить модуль диктофона:\n{e}")
            return
        self.dict_win = DictaphoneWindow(self.root, self.cfg.get("model_repo", MODEL_REPO))
        self.dict_win.focus()

    def refresh_token_badge(self):
        if self.cfg.get("hf_token"):
            self.token_badge.configure(text="● токен задан", text_color="green")
        else:
            self.token_badge.configure(text="● токен не задан", text_color="gray")

    def _toggle_diar(self):
        self.nspk_menu.configure(state="normal" if self.diar_var.get() else "disabled")

    def _lang_hint(self):
        v = self.lang_seg.get()
        if v == "Несколько":
            t = "Смесь русского и английского в одной записи — язык определяется по блокам."
        elif v == "Авто":
            t = "Один язык на всю запись: система сама выберет русский ИЛИ английский (не смесь)."
        else:
            t = f"Вся запись распознаётся как «{v}»."
        self.lang_note.configure(text=t)

    def _open_folder(self):
        if self.last_dir:
            subprocess.run(["open", str(self.last_dir)], check=False)
        else:
            messagebox.showinfo("Папка", "Сначала сделайте хотя бы одну транскрипцию.")

    # -- файлы --
    def _refresh_list(self):
        self.listbox.configure(state="normal")
        self.listbox.delete("1.0", "end")
        if self.files:
            for p in self.files:
                self.listbox.insert("end", "•  " + p.name + "\n")
        else:
            self.listbox.insert("1.0", "Файлы не добавлены")
        self.listbox.configure(state="disabled")

    def _add_paths(self, paths):
        for p in paths:
            path = Path(p)
            if path.suffix.lower() in AUDIO_EXTS and path not in self.files:
                self.files.append(path)
        self._refresh_list()

    def add_files(self):
        self._add_paths(filedialog.askopenfilenames(title="Выберите файлы", filetypes=AUDIO_TYPES))

    def _on_drop(self, event):
        self._add_paths(self.root.tk.splitlist(event.data))

    def clear_files(self):
        self.files.clear(); self._refresh_list()

    # -- лог/таймер --
    def _stat(self, m): self.q.put(("status", m))
    def _prog(self, v): self.q.put(("prog", v))

    def _tick(self):
        if self.running and self.start_time:
            el = int(time.monotonic() - self.start_time)
            self.timer.configure(text=f"⏱ {el // 60:02d}:{el % 60:02d}")
        self.root.after(500, self._tick)

    def _drain(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == "status":
                    self.status.configure(text=payload, text_color="gray")
                elif kind == "prog":
                    if payload < 0:
                        self.progress.configure(mode="indeterminate"); self.progress.start()
                    else:
                        try: self.progress.stop()
                        except Exception: pass
                        self.progress.configure(mode="determinate"); self.progress.set(payload)
                elif kind == "diar_error":
                    messagebox.showwarning(
                        "Говорящие не определены",
                        "Не удалось определить говорящих:\n\n" + payload +
                        "\n\nФайл сохранён без разметки говорящих.")
                elif kind == "done":
                    self._finish(payload)
        except queue.Empty:
            pass
        self.root.after(120, self._drain)

    # -- запуск --
    def start(self):
        if self.running:
            return
        if not self.files:
            messagebox.showinfo("Нет файлов", "Сначала добавьте хотя бы один файл."); return
        if self.diar_var.get() and not self.cfg.get("hf_token"):
            messagebox.showinfo("Нужен токен",
                                "Для определения говорящих задайте токен в «⚙︎ Настройки»."); return
        self.running = True
        self.start_time = time.monotonic()
        self.run_btn.configure(state="disabled")
        self.progress.configure(mode="determinate"); self.progress.set(0)
        sel = self.lang_seg.get()
        params = dict(
            files=list(self.files),
            lang=LANG_CODE.get(sel),
            model=self.cfg.get("model_repo", MODEL_REPO),
            fmt=FMT_SEG[self.fmt_seg.get()],
            with_ts=self.ts_var.get(),
            diar=self.diar_var.get(),
            nspk=NSPK[self.nspk_var.get()],
            token=self.cfg.get("hf_token", ""),
            normalize=self.norm_var.get(),
            multilang=(sel == "Несколько"),
            open_after=True)
        threading.Thread(target=self._worker, kwargs=params, daemon=True).start()

    def _worker(self, files, lang, model, fmt, with_ts, diar, nspk, token, normalize, multilang, open_after):
        try:
            self._stat("Загружаю модель…")
            import mlx_whisper
        except Exception as e:
            self.q.put(("done", ("error", "Не найден mlx_whisper:\n  pip install mlx-whisper\n\n" + str(e)))); return
        if fmt == "docx":
            try:
                import docx  # noqa
            except Exception:
                self.q.put(("done", ("error", "Для Word нужен python-docx:\n  pip install python-docx"))); return
        if diar:
            try:
                import pyannote.audio  # noqa
                import torch  # noqa
            except Exception:
                self.q.put(("done", ("error", "Для говорящих нужен pyannote:\n  pip install pyannote.audio"))); return

        produced = []
        for path in files:
            if not path.exists():
                continue
            self._stat(f"Распознаю: {path.name}")
            duration = ffprobe_duration(path)
            src = path
            tmp_norm = None
            if normalize and not multilang:   # для мультиязыка нормализация вшита в движок
                try:
                    tmp_norm = normalize_audio(path, self._stat)
                    src = tmp_norm
                except Exception as e:
                    self._stat(f"Нормализация не удалась ({e}), обрабатываю как есть.")
            result = None
            try:
                if multilang:
                    self._prog(0)
                    segments = transcribe_multilang_fast(
                        path, model, mlx_whisper, normalize, self._prog, self._stat)
                else:
                    catcher = _ProgressCatcher(duration, self._prog)
                    with contextlib.redirect_stdout(catcher):
                        result = mlx_whisper.transcribe(
                            str(src), path_or_hf_repo=model, language=lang,
                            temperature=0.0, condition_on_previous_text=False, verbose=True)
            except FileNotFoundError as e:
                if "ffmpeg" in str(e):
                    self.q.put(("done", ("error", "Не найден ffmpeg:\n  brew install ffmpeg"))); return
                self._stat(f"Ошибка: {e}"); continue
            except Exception as e:
                self._stat(f"Ошибка на {path.name}: {e}"); continue

            if result is not None:
                segments = result.get("segments", [])
            diar_used = False
            if diar:
                self._prog(-1)
                try:
                    segments = diarize_and_label(path, segments, token, nspk, self._stat, self._prog)
                    diar_used = True
                except Exception as e:
                    self.q.put(("diar_error", str(e)))
                self._prog(1.0)
            out = {"txt": path.with_suffix(".txt"),
                   "docx": path.with_suffix(".docx"),
                   "srt": path.with_suffix(".srt")}[fmt]
            if out.exists():
                try: out.unlink()   # гарантируем перезапись, даже если открыт
                except Exception: pass
            try:
                if fmt == "txt":
                    write_txt(segments, out, with_ts, diar_used)
                elif fmt == "docx":
                    write_docx(segments, out, path.stem, with_ts, diar_used)
                else:
                    write_srt(segments, out, diar_used)
                produced.append(out)
                self.last_dir = out.parent
            except Exception as e:
                self._stat(f"Не удалось сохранить {path.name}: {e}"); continue
            finally:
                if tmp_norm is not None:
                    try: tmp_norm.unlink()
                    except Exception: pass

        if open_after and produced:
            subprocess.run(["open", str(produced[0])], check=False)
        self.q.put(("done", ("ok", f"Готово. Файлов обработано: {len(produced)}.")))

    def _finish(self, payload):
        kind, msg = payload
        self.progress.configure(mode="determinate")
        try: self.progress.stop()
        except Exception: pass
        self.progress.set(1.0 if kind == "ok" else 0)
        self.run_btn.configure(state="normal")
        self.running = False
        el = int(time.monotonic() - self.start_time) if self.start_time else 0
        if kind == "ok":
            self.status.configure(text=f"{msg}  (за {el // 60:02d}:{el % 60:02d})", text_color="green")
        else:
            self.status.configure(text="Ошибка.", text_color="red")
            messagebox.showerror("Ошибка", msg)


class _ProgressCatcher(io.TextIOBase):
    def __init__(self, duration, cb):
        self.duration = duration or 0
        self.cb = cb
        self.buf = ""

    def write(self, s):
        self.buf += s
        if "\n" in self.buf:
            *lines, self.buf = self.buf.split("\n")
            for line in lines:
                m = re.search(r'-->\s*([\d:.]+)\s*\]', line)
                if m and self.duration > 0:
                    self.cb(min(1.0, parse_ts(m.group(1)) / self.duration))
        return len(s)


def main():
    App(RootTk()).root.mainloop()


if __name__ == "__main__":
    main()
