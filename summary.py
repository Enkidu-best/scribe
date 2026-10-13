#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Итоги расшифровки - отдельный документ рядом с полной расшифровкой.

Делается только по запросу (галочка «Итоги» или кнопка «Подвести итоги») и полностью
локально: через Ollama (та же, что у Legal Knowledge Desk), без новых
Python-зависимостей - запросы идут через urllib.

Как работает:
  короткая расшифровка (до ~28 тыс. знаков) - один запрос к модели;
  длинная - сначала каждая часть сжимается в заметки, затем из заметок
  собирается итог («map → reduce»). Итог модель отдаёт строго в JSON по схеме,
  а оформление (Word / текст) делаем сами - так документ всегда одинаковой формы.

Полная расшифровка при этом не меняется: итоги - второй файл
«<имя> - итоги.docx/.txt».
"""
import json
import os
import shutil
import subprocess
import time
import urllib.error
import urllib.request
from pathlib import Path

OLLAMA_URL = (os.environ.get("OLLAMA_HOST") or "http://127.0.0.1:11434").rstrip("/")
if not OLLAMA_URL.startswith("http"):
    OLLAMA_URL = "http://" + OLLAMA_URL
DEFAULT_MODEL = "qwen3.5:9b"
# Если выбранной модели нет - берём первую установленную из этого списка.
PREFERRED = ["qwen3.5:9b", "qwen3.5:4b", "qwen3:8b", "qwen2.5:7b", "llama3.1:8b"]

SINGLE_PASS_CHARS = 28000     # до этого объёма - один запрос
PART_CHARS = 20000            # размер части для длинных записей
NUM_CTX = 16384               # контекст модели: части и заметки помещаются с запасом


class Cancelled(Exception):
    """Пользователь нажал «Остановить»."""


# ---------------------------------------------------------------------------
# Ollama
def _get(path, timeout=3.0):
    with urllib.request.urlopen(OLLAMA_URL + path, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def list_models():
    """Установленные в Ollama модели для текста (без эмбеддингов). Бросает OSError,
    если Ollama не запущена."""
    try:
        data = _get("/api/tags")
    except (urllib.error.URLError, OSError) as e:
        raise OSError(f"Ollama недоступна ({OLLAMA_URL})") from e
    out = []
    for m in data.get("models", []):
        caps = m.get("capabilities")
        if caps is not None and "completion" not in caps:
            continue
        if "embed" in m.get("name", "") or "bge" in m.get("name", ""):
            continue
        out.append(m["name"])
    return out


def ensure_ollama(stat=lambda m: None, wait=25.0):
    """Ollama отвечает? Если нет - пробуем тихо запустить приложение Ollama и ждём."""
    try:
        _get("/api/version", timeout=1.5)
        return True
    except (urllib.error.URLError, OSError):
        pass
    started = False
    if os.path.isdir("/Applications/Ollama.app"):
        subprocess.run(["open", "-g", "-a", "Ollama"], check=False)
        started = True
    elif shutil.which("ollama"):
        subprocess.Popen(["ollama", "serve"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                         start_new_session=True)
        started = True
    if not started:
        return False
    stat("Запускаю Ollama…")
    t0 = time.monotonic()
    while time.monotonic() - t0 < wait:
        time.sleep(1.0)
        try:
            _get("/api/version", timeout=1.5)
            return True
        except (urllib.error.URLError, OSError):
            continue
    return False


def pick_model(wanted, available):
    if wanted and wanted in available:
        return wanted
    for m in PREFERRED:
        if m in available:
            return m
    return available[0] if available else None


def _chat(model, system, user, schema=None, max_tokens=1500, cancel=lambda: False, keep_alive="5m"):
    """Один запрос к /api/chat потоком - чтобы «Остановить» срабатывал сразу,
    а не после конца генерации. keep_alive=0 - модель выгружается сразу после ответа."""
    body = {
        "model": model, "stream": True, "think": False,
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
        "options": {"temperature": 0.2, "num_ctx": NUM_CTX, "num_predict": max_tokens},
        "keep_alive": keep_alive,
    }
    if schema is not None:
        body["format"] = schema
    req = urllib.request.Request(OLLAMA_URL + "/api/chat", data=json.dumps(body).encode("utf-8"),
                                 headers={"Content-Type": "application/json"})
    parts = []
    with urllib.request.urlopen(req, timeout=600) as r:
        for line in r:
            if cancel():
                raise Cancelled()
            if not line.strip():
                continue
            msg = json.loads(line)
            if msg.get("error"):
                raise RuntimeError(msg["error"])
            parts.append(msg.get("message", {}).get("content", ""))
            if msg.get("done"):
                break
    return "".join(parts).strip()


# ---------------------------------------------------------------------------
# Промпты и схема
SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "gist": {"type": "string"},
        "key_points": {"type": "array", "items": {"type": "string"}},
        "decisions": {"type": "array", "items": {"type": "string"}},
        "action_items": {"type": "array", "items": {"type": "string"}},
        "open_questions": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["title", "gist", "key_points", "decisions", "action_items", "open_questions"],
}

_LANG_RULE = ("Пиши по-русски. Если вся запись на английском - пиши по-английски. "
              "Имена, термины и цифры сохраняй как в тексте. Вместо длинного тире пиши короткий дефис «-».")

SYSTEM_FINAL = (
    "Ты подводишь итоги по расшифровке аудиозаписи (встреча, лекция, интервью, разговор). "
    "Опирайся только на текст, ничего не додумывай. Расшифровка сделана автоматически и может "
    "содержать ошибки распознавания - исправляй очевидные, но не выдумывай. " + _LANG_RULE + "\n"
    "Верни JSON:\n"
    "title - короткая тема записи (до 10 слов);\n"
    "gist - суть в 2–4 предложениях;\n"
    "key_points - 4–10 главных фактов и выводов, каждый одной-двумя фразами; сюда НЕ входит то, что попадёт в decisions, action_items или open_questions;\n"
    "decisions - принятые решения и договорённости (пустой список, если их нет);\n"
    "action_items - задачи и следующие шаги: кто - что - срок, если это прозвучало (иначе пустой список);\n"
    "open_questions - нерешённые вопросы (пустой список, если их нет).\n"
    "Всё, о чём договорились или что решили, - в decisions, а не в key_points. "
    "Не повторяй один и тот же пункт в разных разделах.\n"
    "Если в тексте есть метки «Спикер N», указывай, кто что сказал, когда это важно."
)

SYSTEM_PART = (
    "Ты сжимаешь часть длинной расшифровки аудиозаписи в рабочие заметки. "
    "Опирайся только на текст. " + _LANG_RULE + "\n"
    "Выпиши маркированным списком: главные тезисы; решения и договорённости; задачи "
    "(кто - что - срок); нерешённые вопросы. Пропускай пустые разделы. Без вступлений."
)


def transcript_text(segments, diarized):
    """Расшифровка как простой текст: абзацы по паузам, метки говорящих."""
    lines, prev = [], None
    for seg in segments:
        t = (seg.get("text") or "").strip()
        if not t:
            continue
        spk = seg.get("speaker") if diarized else None
        if spk and spk != prev:
            lines.append(f"\n{spk}: {t}")
            prev = spk
        else:
            lines.append(t)
    return " ".join(lines).strip()


def _split(text, size):
    """Режем по границам предложений/абзацев, не посреди слова."""
    parts, rest = [], text
    while len(rest) > size:
        cut = max(rest.rfind("\n", 0, size), rest.rfind(". ", 0, size), rest.rfind("? ", 0, size),
                  rest.rfind("! ", 0, size))
        if cut < size * 0.5:
            cut = rest.rfind(" ", 0, size)
        if cut <= 0:
            cut = size
        parts.append(rest[:cut + 1].strip())
        rest = rest[cut + 1:]
    if rest.strip():
        parts.append(rest.strip())
    return parts


def _clean_list(v):
    return [str(x).strip() for x in (v or []) if str(x).strip()]


def _words(s):
    """Значимые слова по основам (первые 5 букв) - чтобы «согласованию» и «согласовывать» совпали."""
    return {w[:5] for w in "".join(c if c.isalnum() else " " for c in s.lower()).split() if len(w) > 3}


def _dedupe_points(points, others, thr=0.5):
    """Модель любит повторять решения и задачи в «Главном» - убираем такие пункты:
    если большая часть значимых слов пункта уже есть в одном из пунктов других разделов."""
    other_sets = [_words(o) for o in others]
    out = []
    for p in points:
        w = _words(p)
        if w and any(len(w & o) / len(w) >= thr for o in other_sets if o):
            continue
        out.append(p)
    return out or points[:3]


def summarize(text, model, prog=lambda f: None, stat=lambda m: None, cancel=lambda: False):
    """Текст расшифровки -> dict по SCHEMA."""
    if len(text) <= SINGLE_PASS_CHARS:
        source = text
        prog(-1)                     # один запрос - шкала «бежит», без процентов
    else:
        parts = _split(text, PART_CHARS)
        notes = []
        for i, part in enumerate(parts, 1):
            stat(f"часть {i} из {len(parts)}")
            notes.append(f"Часть {i}:\n" + _chat(model, SYSTEM_PART, part, max_tokens=900, cancel=cancel))
            prog(i / (len(parts) + 1))
        source = "Заметки по частям длинной записи (по порядку):\n\n" + "\n\n".join(notes)
        prog(-1)
    stat("сборка итогов")
    # последний запрос - с keep_alive=0: Ollama сразу освобождает память (Mac на 16 ГБ)
    raw = _chat(model, SYSTEM_FINAL, source, schema=SCHEMA, max_tokens=1800, cancel=cancel, keep_alive=0)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        a, b = raw.find("{"), raw.rfind("}")
        data = json.loads(raw[a:b + 1])
    prog(1.0)
    out = {k: _clean_list(data.get(k)) for k in ("key_points", "decisions", "action_items", "open_questions")}
    out["key_points"] = _dedupe_points(out["key_points"],
                                       out["decisions"] + out["action_items"] + out["open_questions"])
    out["title"] = str(data.get("title") or "").strip()
    out["gist"] = str(data.get("gist") or "").strip()
    return out


def _dash(t):
    """Длинное и среднее тире - в короткое: так Андрей пишет сам."""
    return t.replace(" \u2014 ", " - ").replace(" \u2013 ", " - ").replace("\u2014", "-").replace("\u2013", "-")


# ---------------------------------------------------------------------------
# Оформление
_SECTIONS = [("key_points", "Главное"), ("decisions", "Решения и договорённости"),
             ("action_items", "Задачи и следующие шаги"), ("open_questions", "Открытые вопросы")]


def summary_path(transcript: Path) -> Path:
    """Итоги лежат рядом с расшифровкой: Word, если она в Word, иначе .txt."""
    transcript = Path(transcript)
    ext = ".docx" if transcript.suffix.lower() == ".docx" else ".txt"
    return transcript.with_name(f"{transcript.stem} - итоги{ext}")


def _source_line(src_name, transcript_name, duration):
    dur = ""
    if duration:
        m, s = divmod(int(duration), 60)
        h, m = divmod(m, 60)
        dur = f" · {h}:{m:02d}:{s:02d}" if h else f" · {m}:{s:02d}"
    if not src_name:
        return f"Расшифровка: {transcript_name}{dur}"
    return f"Запись: {src_name}{dur} · полная расшифровка: {transcript_name}"


def write_summary(data, path: Path, src_name, transcript_name, duration, model):
    note = (f"Подготовлено автоматически локальной моделью {model} "
            f"{time.strftime('%d.%m.%Y')}. Сверяйте важное с полной расшифровкой.")
    data = {k: ([_dash(x) for x in v] if isinstance(v, list) else _dash(v)) for k, v in data.items()}
    title = data.get("title") or Path(src_name or transcript_name).stem
    tmp = path.with_name("." + path.name + ".tmp")
    if path.suffix == ".docx":
        from docx import Document
        from docx.shared import Pt, RGBColor
        doc = Document()
        doc.add_heading(f"Итоги: {title}", level=1)
        p = doc.add_paragraph()
        r = p.add_run(_source_line(src_name, transcript_name, duration))
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x5f, 0x6b, 0x7a)
        if data.get("gist"):
            doc.add_paragraph(data["gist"])
        for key, head in _SECTIONS:
            items = data.get(key) or []
            if not items:
                continue
            doc.add_heading(head, level=2)
            for it in items:
                doc.add_paragraph(it, style="List Bullet")
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x8a, 0x94, 0xa3)
        doc.save(str(tmp))
    else:
        lines = [f"ИТОГИ: {title}", _source_line(src_name, transcript_name, duration), ""]
        if data.get("gist"):
            lines += [data["gist"], ""]
        for key, head in _SECTIONS:
            items = data.get(key) or []
            if items:
                lines.append(head.upper())
                lines += [f"• {it}" for it in items]
                lines.append("")
        lines.append(note)
        tmp.write_text("\n".join(lines) + "\n", encoding="utf-8")
    os.replace(tmp, path)
    return path
