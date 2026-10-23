#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ядро Scribe без окна: настройки, звук, распознавание («Авто», «Несколько», один язык),
говорящие, запись .txt/.docx/.srt. Его используют окно (whisper_gui.py) и консольный
запуск (scribe_cli.py) - на Mac, Windows и Linux. Распознавание - engine.py.
"""
import json
import os
import subprocess
from pathlib import Path

import platform_util as pu
from summary import Cancelled


# ---------------------------------------------------------------------------
# Turbo - модель по умолчанию (быстрая). Large v3 доступна в Настройках как
# опция «на всякий», но в разы медленнее при том же результате.
MODEL_REPO = "mlx-community/whisper-large-v3-turbo"
MODELS = {
    "Turbo - быстро (рекомендуется)": "mlx-community/whisper-large-v3-turbo",
    "Large v3 - точнее, но в разы медленнее": "mlx-community/whisper-large-v3-mlx",
}
# Язык - сегменты; «Несколько» включает блочный мультиязычный движок.
LANG_SEG = ["Авто", "Русский", "English", "Несколько"]
LANG_CODE = {"Русский": "ru", "English": "en"}   # «Авто»/«Несколько» → None
LANG_HINT = {
    "Авто": "Язык определится сам; если в записи два языка, это тоже будет учтено.",
    "Несколько": "Для записей, где точно перемешаны русский и английский: язык определяется "
                 "для каждой фразы, без предварительной проверки.",
}
LANG_NAME = {"ru": "русская", "en": "английская"}
FMT_SEG = {"Word": "docx", "Текст": "txt", "Субтитры": "srt"}
TEXT_EXTS = {".txt", ".docx", ".srt"}      # готовые расшифровки: из них можно подвести итоги
NSPK = {"Авто": 0, "2": 2, "3": 3, "4": 4, "5": 5, "6": 6}


def _model_title(repo):
    for k, v in MODELS.items():
        if v == repo:
            return k
    return list(MODELS)[0]


def _models_base(cfg):
    return cfg.get("models_dir") or os.environ.get("HF_HOME") \
        or os.path.expanduser("~/.cache/huggingface")


def _models_dir_path(cfg):
    return os.path.join(_models_base(cfg), "hub")


def _models_dir_display(cfg):
    return _models_dir_path(cfg).replace(os.path.expanduser("~"), "~")


def _apply_models_dir(d):
    if d:
        os.environ["HF_HOME"] = d

AUDIO_EXTS = {
    # аудио
    ".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg", ".oga", ".opus",
    ".aiff", ".aif", ".wma", ".m4b", ".amr", ".mka", ".ape", ".wv", ".caf", ".ac3",
    # видео
    ".mp4", ".mov", ".m4v", ".mkv", ".webm", ".wmv", ".avi", ".flv",
    ".3gp", ".3g2", ".ts", ".mts", ".m2ts", ".mpg", ".mpeg", ".asf", ".vob", ".ogv",
}
AUDIO_TYPES = [("Аудио и видео", " ".join("*" + e for e in sorted(AUDIO_EXTS))),
               ("Готовые расшифровки", "*.txt *.docx *.srt"),
               ("Все файлы", "*.*")]

# Этапы обработки файла: строка под шкалой называет этап прямо.
STEP_NAMES = {"prep": "Подготовка звука", "lang": "Разбивка на фразы и определение языка",
              "asr": "Распознавание речи", "diar": "Определение говорящих", "sum": "Итоги"}
STEP_CHIP = {"prep": "Подготовка", "lang": "Языки", "asr": "Распознавание", "diar": "Говорящие",
             "sum": "Итоги", "save": "Сохранение"}
DIAR_STEPS = {"segmentation": "разбивка на фрагменты", "speaker_counting": "подсчёт говорящих",
              "embeddings": "голосовые отпечатки", "discrete_diarization": "группировка"}

# Список поддерживаемых форматов - показываем в «Настройках».
FORMATS_AUDIO = "MP3, WAV, M4A, AAC, FLAC, OGG, OPUS, AIFF, WMA, M4B, AMR"
FORMATS_VIDEO = "MP4, MOV, MKV, WEBM, WMV, AVI, FLV, 3GP, TS, MPG"

APP_VERSION = "2.7"
CONFIG_PATH = pu.config_path()     # Mac - прежний ~/.audio2text.json; Windows/Linux - папка настроек системы
# Пробуем модели по порядку: сначала новая community-1 (её тянут свежие версии
# pyannote и она точнее), затем классическая 3.1 как запасная.
DIAR_MODELS = [
    "pyannote/speaker-diarization-community-1",
    "pyannote/speaker-diarization-3.1",
]
# Что проверять на доступ: community-1 самодостаточна; 3.1 требует ещё segmentation-3.0.
DIAR_CHECK = [
    "pyannote/speaker-diarization-community-1",
    "pyannote/speaker-diarization-3.1",
    "pyannote/segmentation-3.0",
]


def load_config() -> dict:
    try:
        return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_config(cfg: dict) -> None:
    try:
        CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
        CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


# ---------------------------------------------------------------------------
def ffprobe_duration(path) -> float:
    try:
        out = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
            capture_output=True, text=True)
        return float(out.stdout.strip())
    except Exception:
        return 0.0


def _clock(seconds: float, srt: bool = False) -> str:
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    if srt:
        return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def group_paragraphs(segments, gap=1.2, max_chars=380):
    paras, buf, start, prev_end = [], [], None, None
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        if start is None:
            start = seg["start"]
        big_pause = prev_end is not None and (seg["start"] - prev_end) > gap
        cur = " ".join(buf)
        long_enough = len(cur) > max_chars and cur.rstrip().endswith((".", "!", "?", "…"))
        if buf and (big_pause or long_enough):
            paras.append((start, " ".join(buf).strip()))
            buf, start = [], seg["start"]
        buf.append(text)
        prev_end = seg["end"]
    if buf:
        paras.append((start, " ".join(buf).strip()))
    return paras


def group_by_speaker(segments, gap=1.2, max_chars=380):
    """Блоки (start, speaker, text): новый блок при смене говорящего,
    большой паузе или когда абзац стал длинным."""
    blocks, buf, start, cur, prev_end = [], [], None, None, None
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        spk = seg.get("speaker", "Спикер ?")
        if cur is None:
            cur, start = spk, seg["start"]
        speaker_changed = spk != cur
        big_pause = prev_end is not None and (seg["start"] - prev_end) > gap
        cur_text = " ".join(buf)
        long_enough = len(cur_text) > max_chars and cur_text.rstrip().endswith((".", "!", "?", "…"))
        if buf and (speaker_changed or big_pause or long_enough):
            blocks.append((start, cur, " ".join(buf).strip()))
            buf, start, cur = [], seg["start"], spk
        buf.append(text)
        prev_end = seg["end"]
    if buf:
        blocks.append((start, cur, " ".join(buf).strip()))
    return blocks


# Запись - во временный файл рядом и затем os.replace: результат перезаписывается
# даже если старый файл открыт, а при сбое старый файл не теряется.
def _tmp_for(path):
    return path.with_name("." + path.name + ".tmp")


def write_txt(segments, path, with_ts, diarized):
    lines = []
    if diarized:
        prev_spk = None
        for start, spk, text in group_by_speaker(segments):
            ts = f"[{_clock(start)}] " if with_ts else ""
            if spk != prev_spk:
                lines.append(f"{ts}{spk}: {text}")
            else:
                lines.append(f"{ts}{text}")
            lines.append("")
            prev_spk = spk
    else:
        for start, text in group_paragraphs(segments):
            p = f"[{_clock(start)}] " if with_ts else ""
            lines.append(f"{p}{text}"); lines.append("")
    tmp = _tmp_for(path)
    tmp.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    os.replace(tmp, path)


def write_srt(segments, path, diarized):
    blocks = []
    for i, seg in enumerate(segments, 1):
        text = seg["text"].strip()
        if diarized and seg.get("speaker"):
            text = f"{seg['speaker']}: {text}"
        blocks.append(str(i))
        blocks.append(f"{_clock(seg['start'], True)} --> {_clock(seg['end'], True)}")
        blocks.append(text); blocks.append("")
    tmp = _tmp_for(path)
    tmp.write_text("\n".join(blocks), encoding="utf-8")
    os.replace(tmp, path)


def write_docx(segments, path, title, with_ts, diarized):
    from docx import Document
    from docx.shared import RGBColor
    doc = Document()
    doc.add_heading(title, level=1)
    if diarized:
        prev_spk = None
        for start, spk, text in group_by_speaker(segments):
            p = doc.add_paragraph()
            if with_ts:
                ts = p.add_run(f"[{_clock(start)}] "); ts.italic = True
                ts.font.color.rgb = RGBColor(0x5f, 0x6b, 0x7a)
            if spk != prev_spk:
                r = p.add_run(f"{spk}: "); r.bold = True
                r.font.color.rgb = RGBColor(0x1f, 0x38, 0x64)
            p.add_run(text)
            prev_spk = spk
    else:
        for start, text in group_paragraphs(segments):
            p = doc.add_paragraph()
            if with_ts:
                t = p.add_run(f"[{_clock(start)}] "); t.italic = True
                t.font.color.rgb = RGBColor(0x5f, 0x6b, 0x7a)
            p.add_run(text)
    tmp = _tmp_for(path)
    doc.save(str(tmp))
    os.replace(tmp, path)


# ---------------------------------------------------------------------------
# Звук: один проход ffmpeg в память (и нормализация громкости в том же проходе).
SR_FAST = 16000


# --- ошибки: по-русски и с действием; исходный текст - в журнал (стандарт 2.3, Scribe 2.7) ---
_MEM_SIGNS = ("out of memory", "outofmemory", "insufficient memory", "metal::malloc", "kiogpucommandbuffer",
              "memoryerror", "cannot allocate", "failed to allocate")


def heavy_lock_owner():
    """Кто держит общий замок видеокарты (../.heavy-lock рядом с папкой проекта): «Herald» или None."""
    try:
        owner = (pu.source_dir().parent / ".heavy-lock" / "owner").read_text(encoding="utf-8").strip()
    except Exception:
        return None
    if not owner or owner.startswith("Scribe"):
        return None
    words = owner.split()
    i = next((k for k, w in enumerate(words) if w[:3] in ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")), len(words))
    return " ".join(words[:i]) or None


def is_memory_error(e):
    t = f"{type(e).__name__} {e}".lower()
    return isinstance(e, MemoryError) or any(x in t for x in _MEM_SIGNS)


def explain_error(e, resume=False):
    """Текст ошибки для окна: что случилось и что сделать. Английский - только в журнал."""
    t = str(e)
    again = "Нажмите «Расшифровать» ещё раз" + (" - продолжу с места остановки." if resume else ".")
    if "no Stream" in t or "Stream(gpu" in t:
        return "Сбой движка распознавания. " + again
    if is_memory_error(e):
        who = heavy_lock_owner()
        return ("Не хватило памяти видеокарты." + (f" Сейчас её занимает «{who}»." if who else "")
                + " Подождите, пока она освободится, и " + again[0].lower() + again[1:])
    if t.startswith("ffmpeg:"):
        return "Не удалось прочитать файл - возможно, он повреждён или это не аудио и не видео."
    if isinstance(e, RuntimeError) and any("а" <= c.lower() <= "я" for c in t[:40]):
        return t.strip().splitlines()[0]          # наше сообщение - уже по-русски
    return "Ошибка при обработке файла. " + again + " Подробности - в журнале Scribe."


def log_error(e, where=""):
    """Полный текст и трассировка - в журнал (logs/scribe-errors.log)."""
    import traceback
    import time as _t
    try:
        d = pu.log_dir()
        d.mkdir(parents=True, exist_ok=True)
        with open(d / "scribe-errors.log", "a", encoding="utf-8") as f:
            f.write(f"\n--- {_t.strftime('%Y-%m-%d %H:%M:%S')} {where}\n")
            f.write("".join(traceback.format_exception(type(e), e, e.__traceback__)))
    except Exception:
        pass


def _decode_to_array(path, normalize=True, cancel=None, prog=None, duration=None):
    """ОДИН проход ffmpeg -> float32 моно 16 кГц в память. loudnorm вшит сюда же,
    так что отдельного прохода нормализации не нужно. Читаем по кускам, чтобы
    «Остановить» прерывал и долгую подготовку звука.
    prog(доля) - процент этапа (2.7): ffmpeg отдаёт звук потоком, полученные секунды против
    длительности файла и есть прогресс. Звук копится в одном bytearray, без лишних копий:
    на 3-часовой записи это ~690 МБ вместо ~2 ГБ в пике."""
    import numpy as np
    import tempfile
    cmd = ["ffmpeg", "-nostdin", "-i", str(path)]
    if normalize:
        cmd += ["-af", "loudnorm"]
    cmd += ["-ar", str(SR_FAST), "-ac", "1", "-f", "f32le", "-acodec", "pcm_f32le",
            "-hide_banner", "-loglevel", "error", "-"]
    if prog is not None and not duration:
        duration = ffprobe_duration(path)
    total = duration * SR_FAST * 4 if duration else 0
    with tempfile.TemporaryFile() as err:
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=err)
        data = bytearray()
        last = 0.0
        while True:
            if cancel is not None and cancel():
                proc.kill()
                proc.wait()
                raise Cancelled()
            buf = proc.stdout.read(1 << 20)
            if not buf:
                break
            data += buf
            if prog is not None and total:
                f = min(0.99, len(data) / total)
                if f - last >= 0.01:
                    last = f
                    prog(f)
        proc.wait()
        if proc.returncode != 0:
            err.seek(0)
            raise RuntimeError("ffmpeg: " + err.read().decode("utf-8", "ignore")[-300:])
    n = len(data) // 4 * 4
    return np.frombuffer(memoryview(data)[:n], dtype=np.float32)


def _transcribe_array(eng, audio, model_repo, lang, prog, cancel, initial_prompt=None, temperature=0.0):
    """Один вызов движка распознавания (engine.py: MLX или faster-whisper) с прогрессом и отменой."""
    return eng.transcribe(audio, model_repo, lang, initial_prompt, temperature, prog, cancel)


def transcribe_auto(audio, model_repo, eng, prog, cancel, stage=None, forced=None, on_mixed=None):
    """«Авто» (2.3): сначала дешёвая проверка языков по всем фразам (multilang.probe_languages:
    сито whisper-base + перепроверка turbo).
    В записи есть и русский, и английский - обрабатываем как «Несколько»; иначе - один язык,
    найденный проверкой (а не только по первым 30 с, как решал бы сам Whisper).
    forced - язык, выбранный вручную: проверки языков нет вообще (2.5, решение Андрея) - пользователь
    выбрал язык сам, вся запись распознаётся на нём.
    -> (segments, info)"""
    import multilang
    stage = stage or (lambda k, d="": None)
    if forced:
        info = dict(mode="single", lang=forced)
        stage("asr")
        prog(0)
        return _transcribe_array(eng, audio, model_repo, forced, prog, cancel), info
    stage("prep", "проверка языков")
    probe = multilang.probe_languages(audio, model_repo, cancel, prog)
    info = dict(probe=probe)
    if probe["mixed"]:
        if on_mixed:
            on_mixed(probe)
        info["mode"] = "mixed"
        segs = transcribe_multilang_fast(audio, model_repo, eng, prog, lambda m: None, cancel, stage=stage)
        return segs, info
    else:
        # язык, найденный проверкой по всей записи (а не только по первым 30 с, как решил бы
        # сам Whisper) - если проверка в нём уверена
        lang = probe["dominant"] if probe["dominant_share"] >= 0.8 and probe["phrases"] >= 3 else None
    info.update(mode="single", lang=lang)
    stage("asr")
    prog(0)
    return _transcribe_array(eng, audio, model_repo, lang, prog, cancel), info


def transcribe_multilang_fast(audio, model_repo, eng, prog, stat, cancel, stage=None, stats=None):
    """Режим «Несколько»: язык - для каждой фразы (см. multilang.py)."""
    import multilang

    def block(clip, lang, prompt, bprog, temperature=0.0):
        return _transcribe_array(eng, clip, model_repo, lang, bprog, cancel, initial_prompt=prompt,
                                 temperature=temperature)

    return multilang.transcribe(audio, model_repo, block, prog, cancel,
                                stage=stage or (lambda k, d="": None), stats=stats)


# ---------------------------------------------------------------------------
_PIPELINES = {}   # модель диаризации грузится один раз на запуск приложения


def _load_pipeline(token):
    """Загружает pyannote pipeline независимо от версии: логинимся заранее,
    токен подхватывается из кэша/окружения, а не через параметр (у разных
    версий pyannote параметр называется по-разному: token / use_auth_token)."""
    if token in _PIPELINES:
        return _PIPELINES[token]
    if token:                      # без токена - только если модель уже на диске (офлайн-пакет)
        os.environ["HF_TOKEN"] = token
        os.environ["HUGGING_FACE_HUB_TOKEN"] = token
        try:
            from huggingface_hub import login
            login(token=token, add_to_git_credential=False)
        except Exception:
            pass
    from pyannote.audio import Pipeline
    errors = []
    if not token:                  # офлайн-пакет: грузим из папки, не обращаясь к Hugging Face
        try:
            from huggingface_hub import snapshot_download
            pl = Pipeline.from_pretrained(snapshot_download(DIAR_MODELS[0], local_files_only=True))
            if pl is not None:
                _PIPELINES[token] = pl
                return pl
        except Exception as e:
            errors.append(f"{DIAR_MODELS[0]} (с диска): {e}")
    for model in DIAR_MODELS:
        for kw in ({}, {"token": token}):
            try:
                pl = Pipeline.from_pretrained(model, **kw)
                if pl is not None:
                    _PIPELINES[token] = pl
                    return pl
            except Exception as e:
                errors.append(f"{model}: {e}")
    raise RuntimeError(
        "Ни одна модель диаризации не загрузилась. Обычно причина - не приняты "
        "условия модели на Hugging Face. Откройте и нажмите Agree / Accept на:\n"
        "huggingface.co/pyannote/speaker-diarization-community-1\n\n"
        + "\n".join(errors[:4]))


def _mk_hook(prog, stat, cancel):
    """Колбэк прогресса pyannote: шаги диаризации - в шкалу и строку этапа (по-русски),
    заодно точка отмены."""
    last = [None]

    def hook(step_name, step_artifact=None, file=None, total=None, completed=None):
        if cancel():
            raise Cancelled()
        try:
            if step_name and step_name != last[0]:
                last[0] = step_name
                stat(DIAR_STEPS.get(step_name, "обработка"))
            if total and completed is not None and float(total) > 0:
                prog(min(1.0, float(completed) / float(total)))
        except Exception:
            pass
    return hook


def _extract_turns(diar):
    """Достаёт (start, end, speaker) из результата любой версии pyannote."""
    ann = getattr(diar, "speaker_diarization", diar)
    turns = []
    if hasattr(ann, "itertracks"):
        for item in ann.itertracks(yield_label=True):
            seg, spk = item[0], item[-1]
            turns.append((seg.start, seg.end, spk))
    else:
        for turn, spk in ann:
            turns.append((turn.start, turn.end, spk))
    return turns


def diarize_and_label(audio_path, segments, token, num_speakers, log, prog, cancel=lambda: False):
    import torch
    log("загрузка модели")
    pipeline = _load_pipeline(token)
    device = "cpu"
    try:
        if torch.backends.mps.is_available():
            pipeline.to(torch.device("mps")); device = "mps"
        elif torch.cuda.is_available():
            pipeline.to(torch.device("cuda")); device = "cuda"
    except Exception:
        device = "cpu"
    # Готовим аудио сами через ffmpeg и подаём waveform напрямую - минуя torchcodec,
    # который ломается на свежих связках PyTorch/ffmpeg. Без нормализации - как раньше.
    data = _decode_to_array(audio_path, normalize=False)
    waveform = torch.from_numpy(data).unsqueeze(0)  # (1, N)
    payload = {"waveform": waveform, "sample_rate": SR_FAST}
    kwargs = {"num_speakers": num_speakers} if num_speakers and num_speakers > 0 else {}
    hook = _mk_hook(prog, log, cancel)

    def _run(pl):
        try:
            return pl(payload, hook=hook, **kwargs)
        except TypeError:
            return pl(payload, **kwargs)  # старые версии без hook

    try:
        diar = _run(pipeline)
    except Cancelled:
        raise
    except Exception:
        if device == "cpu":
            raise
        log("повтор на процессоре")
        pipeline.to(torch.device("cpu"))
        diar = _run(pipeline)
    turns = _extract_turns(diar)
    mapping, n = {}, 0
    for _, _, spk in turns:
        if spk not in mapping:
            n += 1; mapping[spk] = f"Спикер {n}"

    def speaker_at(a, b):
        best, best_ov = None, 0.0
        for s, e, spk in turns:
            ov = max(0.0, min(b, e) - max(a, s))
            if ov > best_ov:
                best_ov, best = ov, spk
        return mapping.get(best, "Спикер ?")

    for seg in segments:
        seg["speaker"] = speaker_at(seg["start"], seg["end"])
    return segments


def free_asr_models():
    """Перед итогами выгружаем Whisper и pyannote: qwen3.5:9b занимает ~6 ГБ, и вместе
    с ними (и с работающим LKD) Mac на 16 ГБ начинает задыхаться."""
    import gc
    _PIPELINES.clear()
    try:
        import multilang
        multilang._SIEVE.clear()
    except Exception:
        pass
    try:
        import engine
        engine.get_engine().unload()
    except Exception:
        pass
    gc.collect()
    try:
        import torch
        if torch.backends.mps.is_available():
            torch.mps.empty_cache()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    except Exception:
        pass


def read_transcript(path):
    """Текст готовой расшифровки (.txt / .docx / .srt) - для итогов по файлам из прошлых запусков."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document
        paras = [p.text for p in Document(str(path)).paragraphs]
        if paras and paras[0].strip() == path.stem:      # заголовок, который пишет сам Scribe
            paras = paras[1:]
        return "\n".join(t for t in paras if t.strip())
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".srt":
        keep = [ln for ln in text.splitlines()
                if ln.strip() and not ln.strip().isdigit() and "-->" not in ln]
        return " ".join(keep)
    return text


def ollama_installed():
    return pu.ollama_installed()


def _fmt_dur(sec):
    if not sec:
        return ""
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def _fmt_size(n):
    for unit in ("Б", "КБ", "МБ", "ГБ"):
        if n < 1024 or unit == "ГБ":
            return f"{n:.0f} {unit}" if unit in ("Б", "КБ") else f"{n:.1f} {unit}".replace(".", ",")
        n /= 1024


def _short(name, n=58):
    return name if len(name) <= n else name[:n - 12] + "…" + name[-11:]
